# from fastapi import APIRouter, Depends
# from fastapi.responses import FileResponse
# from sqlalchemy.orm import Session
# from openpyxl import Workbook

# from dependencies import get_db
# from models import Expense, User
# from auth import get_current_user

# from openpyxl.drawing.image import Image as ExcelImage
# import os

# router = APIRouter()


# @router.get("/export-excel")
# def export_excel(
#     db: Session = Depends(get_db),
#     user: User = Depends(get_current_user)
# ):

#     print("export excel api start")

#     expenses = db.query(Expense).filter(
#         Expense.user_id == user.id
#     ).all()

#     # ================= TERMINAL PRINT START =================
#     print("\n================ ALL EXPENSES ================\n")

#     wb = Workbook()

#     ws1 = wb.active
#     ws1.title = "All Expenses"

#     ws1.append(["Title", "Amount", "Type", "Category", "Description", "Date"])

#     for exp in expenses:

#         # 👉 PRINT ADDED (same data as Excel)
#         print(
#     f"Title: {exp.title} | "
#     f"Amount: {exp.amount} | "
#     f"Type: {exp.type} | "
#     f"Category: {exp.category} | "
#     f"Description: {exp.description} | "
#     f"Date: {exp.date.strftime('%Y-%m-%d')}"
# )

#         ws1.append([
#             exp.title,
#             exp.amount,
#             exp.type,
#             exp.category,
#             exp.description,
#             exp.date.strftime("%Y-%m-%d")
#         ])

#     # ================= DAY WISE =================
#     ws2 = wb.create_sheet("Day Wise")
#     print("\n================ DAY WISE ================\n")

#     day_data = {}
#     for exp in expenses:
#         day = exp.date.strftime("%Y-%m-%d")
#         day_data[day] = day_data.get(day, 0) + exp.amount

#     ws2.append(["Date", "Total Amount"])

#     for day, total in day_data.items():
#         print(day, "=>", total)
#         ws2.append([day, total])

#     # ================= MONTHLY =================
#     ws3 = wb.create_sheet("Monthly")
#     print("\n================ MONTHLY ================\n")

#     month_data = {}
#     for exp in expenses:
#         month = exp.date.strftime("%Y-%m")
#         month_data[month] = month_data.get(month, 0) + exp.amount

#     ws3.append(["Month", "Total Amount"])

#     for month, total in month_data.items():
#         print(month, "=>", total)
#         ws3.append([month, total])

#     # ================= CATEGORY WISE =================
#     ws4 = wb.create_sheet("Category Wise")
#     print("\n================ CATEGORY WISE ================\n")

#     category_data = {}
#     for exp in expenses:
#         category_data[exp.category] = category_data.get(exp.category, 0) + exp.amount

#     ws4.append(["Category", "Total Amount"])

#     for category, total in category_data.items():
#         print(category, "=>", total)
#         ws4.append([category, total])

#     file_path = "expenses_report.xlsx"
#     wb.save(file_path)

#     return FileResponse(
#         path=file_path,
#         filename="expenses_report.xlsx",
#         media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#     )

from fastapi import APIRouter, Depends
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook

from dependencies import get_db
from models import Expense, User
from auth import get_current_user

from openpyxl.drawing.image import Image as ExcelImage
import os

router = APIRouter()
BASE_DIR = os.path.dirname(os.path.abspath(__file__))


@router.get("/export-excel")
def export_excel(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user)
):

    print("export excel api start")

    expenses = db.query(Expense).filter(
        Expense.user_id == user.id
    ).all()

    print("\n================ ALL EXPENSES ================\n")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "All Expenses"

    # ✅ IMAGE COLUMN ADDED
    ws1.append(["Title", "Amount", "Type", "Category", "Description", "Date", "Image"])

    row = 2  # ✅ for image placement

    for exp in expenses:

        print(
            f"Title: {exp.title} | "
            f"Amount: {exp.amount} | "
            f"Type: {exp.type} | "
            f"Category: {exp.category} | "
            f"Description: {exp.description} | "
            f"Date: {exp.date.strftime('%Y-%m-%d')}"
        )

        ws1.append([
            exp.title,
            exp.amount,
            exp.type,
            exp.category,
            exp.description,
            exp.date.strftime("%Y-%m-%d")
        ])

        # ================= IMAGE INSERT =================
        if exp.image:
          image_path = os.path.join(os.getcwd(), exp.image)
          image_path = os.path.abspath(image_path)

          print("IMAGE PATH:", image_path)
          print("DB IMAGE:", exp.image)
          print("FULL IMAGE PATH:", image_path)
          print("EXISTS:", os.path.exists(image_path))

          if image_path and os.path.exists(image_path):
              try:
                  img = ExcelImage(image_path)
                  img.width = 90
                  img.height = 90
                  ws1.row_dimensions[row].height = 70
                  ws1.column_dimensions["G"].width = 15

                # 🔥 FIXED COLUMN G + ROW MATCH
                  ws1.add_image(img, f"G{row}")
                  print("images added")

              except Exception as e:
                  print("Image error:", e)

        row += 1

    # ================= DAY WISE =================
    ws2 = wb.create_sheet("Day Wise")
    print("\n================ DAY WISE ================\n")

    day_data = {}
    for exp in expenses:
        day = exp.date.strftime("%Y-%m-%d")
        day_data[day] = day_data.get(day, 0) + exp.amount

    ws2.append(["Date", "Total Amount"])

    for day, total in day_data.items():
        print(day, "=>", total)
        ws2.append([day, total])

    # ================= MONTHLY =================
    ws3 = wb.create_sheet("Monthly")
    print("\n================ MONTHLY ================\n")

    month_data = {}
    for exp in expenses:
        month = exp.date.strftime("%Y-%m")
        month_data[month] = month_data.get(month, 0) + exp.amount

    ws3.append(["Month", "Total Amount"])

    for month, total in month_data.items():
        print(month, "=>", total)
        ws3.append([month, total])

    # ================= CATEGORY WISE =================
    ws4 = wb.create_sheet("Category Wise")
    print("\n================ CATEGORY WISE ================\n")

    category_data = {}
    for exp in expenses:
        category_data[exp.category] = category_data.get(exp.category, 0) + exp.amount

    ws4.append(["Category", "Total Amount"])

    for category, total in category_data.items():
        print(category, "=>", total)
        ws4.append([category, total])

    # ================= SAVE FILE =================
    file_path = f"expenses_report_{user.id}.xlsx"
    wb.save(file_path)

    return FileResponse(
        path=file_path,
        filename="expenses_report.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )